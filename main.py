"""
Migrate data from an existing (old) Excel file into a new Excel file whose
column names are different, using the mapping in data_migration_helper.xlsx.

Every row of the helper's "Mapping" sheet becomes one column of the new file,
in the same order. The column is filled with an exact copy of the data under
the matching "Old Column" of the old file: values and their number formats,
so dates stay dates, codes keep their leading zeros, decimals are unchanged.

Usage:
    python main.py --old old_data.xlsx --output new_data.xlsx
    python main.py          (no options: asks for everything step by step)
    python main.py --help   (explains every option)

Options:
    --helper       mapping workbook (default: data_migration_helper.xlsx next to this script)
    --sheet        sheet of the old file to read (required only if it has several sheets)
    --header-row   row number that holds the column names in the old file (default: 1)
    --new-sheet    sheet name for the new file (default: Sheet1)
    --overwrite    replace the output file if it already exists

Safety:
    * The old file and the helper are only read, never written.
    * The new file is first written to "<output>.incomplete.xlsx", re-read and
      compared cell by cell with the old file, and only renamed to the output
      path when every cell matches.
    * Messages show cell positions and column names only, never cell contents.
"""

from __future__ import annotations

import argparse
import difflib
import os
import sys
import warnings
import zipfile
from dataclasses import dataclass
from itertools import zip_longest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.read_only import EMPTY_CELL
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

MAPPING_SHEET = "Mapping"
NEW_COL_HEADER = "New Column"
OLD_COL_HEADER = "Old Column"
MAX_LISTED = 20  # cap on cell positions listed in one error message


class MigrationError(Exception):
    pass


@dataclass
class ColumnMap:
    helper_row: int
    new_name: str
    old_name: str | None  # None: the new column is created empty
    old_index: int | None = None  # 0-based column position in the old sheet


def is_blank(value):
    return value is None or value == ""


def text(value):
    return "" if value is None else str(value).strip()


def normalize(name):
    """Key for comparing column names: ignores upper/lower case and extra spaces."""
    return " ".join(str(name).split()).casefold()


def at(row, index, default=None):
    """row[index], or default when the row is shorter (Excel omits trailing empty cells)."""
    return row[index] if index < len(row) else default


def cell_ref(col_index, row_no):
    return f"{get_column_letter(col_index + 1)}{row_no}"


def listed(items):
    items = list(items)
    shown = ", ".join(items[:MAX_LISTED])
    return shown + (f" ... and {len(items) - MAX_LISTED} more" if len(items) > MAX_LISTED else "")


def open_workbook(path, data_only=True):
    try:
        return load_workbook(path, read_only=True, data_only=data_only)
    except (InvalidFileException, zipfile.BadZipFile, KeyError) as exc:
        raise MigrationError(
            f"cannot open {path} as an .xlsx file ({exc}). "
            "If it is an old .xls file, open it in Excel and 'Save As' .xlsx first."
        ) from None


def pick_sheet(wb, path, sheet_name):
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise MigrationError(
                f"{path} has no sheet named '{sheet_name}'. Sheets: {', '.join(wb.sheetnames)}"
            )
        ws = wb[sheet_name]
    elif len(wb.sheetnames) == 1:
        ws = wb.worksheets[0]
    else:
        raise MigrationError(
            f"{path} has {len(wb.sheetnames)} sheets ({', '.join(wb.sheetnames)}); "
            "choose one with --sheet."
        )
    # Some programs store a wrong sheet size, which would make openpyxl stop early.
    # Forget it so every row actually present in the file is read.
    ws.reset_dimensions()
    return ws


def read_mapping(helper_path):
    wb = open_workbook(helper_path)
    try:
        if MAPPING_SHEET not in wb.sheetnames:
            raise MigrationError(f"{helper_path} has no sheet named '{MAPPING_SHEET}'.")
        ws = wb[MAPPING_SHEET]
        ws.reset_dimensions()
        rows = ws.iter_rows(values_only=True)
        header = [normalize(text(v)) for v in next(rows, ())]
        try:
            new_i = header.index(normalize(NEW_COL_HEADER))
            old_i = header.index(normalize(OLD_COL_HEADER))
        except ValueError:
            raise MigrationError(
                f"row 1 of the '{MAPPING_SHEET}' sheet in {helper_path} must contain "
                f"'{NEW_COL_HEADER}' and '{OLD_COL_HEADER}'."
            ) from None

        mappings, errors, used = [], [], {}
        for row_no, row in enumerate(rows, start=2):
            new_name, old_name = text(at(row, new_i)), text(at(row, old_i))
            if not new_name and not old_name:
                continue
            if not new_name:
                errors.append(f"row {row_no}: '{OLD_COL_HEADER}' is filled but '{NEW_COL_HEADER}' is empty.")
                continue
            key = normalize(new_name)
            if key in used:
                errors.append(f"row {row_no}: new column '{new_name}' is already used in row {used[key]}.")
                continue
            used[key] = row_no
            mappings.append(ColumnMap(row_no, new_name, old_name or None))
    finally:
        wb.close()

    if errors:
        raise MigrationError(f"problems in {helper_path}:\n  " + "\n  ".join(errors))
    if not any(m.old_name for m in mappings):
        raise MigrationError(f"{helper_path} does not map any old column; nothing to migrate.")
    return mappings


def resolve_old_columns(mappings, old_header, header_row):
    """Find the position of every mapped old column in the old file's header row."""
    positions, display = {}, {}
    for index, value in enumerate(old_header):
        if not text(value):
            continue
        key = normalize(value)
        positions.setdefault(key, []).append(index)
        display.setdefault(key, text(value))

    errors = []
    for m in mappings:
        if m.old_name is None:
            continue
        found = positions.get(normalize(m.old_name), [])
        if len(found) == 1:
            m.old_index = found[0]
        elif not found:
            close = difflib.get_close_matches(normalize(m.old_name), list(positions), n=1)
            hint = f" Did you mean '{display[close[0]]}'?" if close else ""
            errors.append(
                f"helper row {m.helper_row}: old column '{m.old_name}' is not in row {header_row} "
                f"of the old file.{hint}"
            )
        else:
            columns = ", ".join(get_column_letter(i + 1) for i in found)
            errors.append(
                f"helper row {m.helper_row}: old column '{m.old_name}' appears more than once "
                f"in the old file (columns {columns}); rename one of them."
            )
    if errors:
        raise MigrationError("column mapping does not match the old file:\n  " + "\n  ".join(errors))

    mapped = {m.old_index for m in mappings}
    return [display[normalize(v)] for i, v in enumerate(old_header) if text(v) and i not in mapped]


def copy_cell(out_ws, src):
    """Exact copy of a cell's value and number format."""
    if is_blank(src.value):
        return None
    dst = WriteOnlyCell(out_ws, value=src.value)
    if isinstance(src.value, str):
        # Keep text as text: without this, text like "=abc" would be written as a formula.
        dst.data_type = "e" if src.data_type == "e" else "s"
    dst.number_format = src.number_format
    return dst


def write_new_file(old_path, sheet_name, header_row, mappings, out_ws):
    """Copy the mapped columns into out_ws.

    Returns (rows written, formula cells copied, old columns not migrated, unnamed columns with data).
    """
    values_wb = open_workbook(old_path, data_only=True)  # what Excel shows (formula results)
    formulas_wb = open_workbook(old_path, data_only=False)  # to detect formulas
    try:
        values_ws = pick_sheet(values_wb, old_path, sheet_name)
        formulas_ws = pick_sheet(formulas_wb, old_path, sheet_name)
        old_header = next(values_ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
        if not any(text(v) for v in old_header):
            raise MigrationError(f"row {header_row} of {old_path} is empty; set the right --header-row.")
        not_migrated = resolve_old_columns(mappings, old_header, header_row)

        header_cells = []
        for m in mappings:
            cell = WriteOnlyCell(out_ws, value=m.new_name)
            cell.data_type = "s"
            header_cells.append(cell)
        out_ws.append(header_cells)

        named = {i for i, v in enumerate(old_header) if text(v)}
        unnamed_with_data = set()
        formula_cells, formulas_without_result = 0, []
        rows_written, pending_blank_rows = 0, 0

        rows = zip(values_ws.iter_rows(min_row=header_row + 1), formulas_ws.iter_rows(min_row=header_row + 1))
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            for row_no, (vrow, frow) in enumerate(rows, start=header_row + 1):
                filled = [i for i, c in enumerate(vrow) if not is_blank(c.value)]
                if not filled:
                    # Blank rows inside the data are kept so old row N stays new row N;
                    # blank rows at the very end are dropped.
                    pending_blank_rows += 1
                    continue
                unnamed_with_data.update(i for i in filled if i not in named)
                for _ in range(pending_blank_rows):
                    out_ws.append([])
                rows_written += pending_blank_rows + 1
                pending_blank_rows = 0

                out_row = []
                for m in mappings:
                    if m.old_index is None:
                        out_row.append(None)
                        continue
                    src = at(vrow, m.old_index, EMPTY_CELL)
                    if at(frow, m.old_index, EMPTY_CELL).data_type == "f":
                        formula_cells += 1
                        if src.value is None:
                            formulas_without_result.append(cell_ref(m.old_index, row_no))
                    out_row.append(copy_cell(out_ws, src))
                out_ws.append(out_row)
            bad_dates = [str(w.message) for w in caught if "outside the limits for dates" in str(w.message)]
    finally:
        values_wb.close()
        formulas_wb.close()

    if formulas_without_result:
        raise MigrationError(
            "these cells of the old file contain formulas whose result was never saved, "
            f"so their value cannot be read: {listed(formulas_without_result)}.\n"
            "  Open the old file in Excel, save it, and run again (or replace the formulas with values)."
        )
    if bad_dates:
        raise MigrationError("some date cells in the old file cannot be read exactly:\n  " + "\n  ".join(bad_dates))
    unnamed = sorted(get_column_letter(i + 1) for i in unnamed_with_data)
    return rows_written, formula_cells, not_migrated, unnamed


def same_cell(src, out):
    if is_blank(src.value) and is_blank(out.value):
        return True
    return (
        type(src.value) is type(out.value)
        and src.value == out.value
        and src.data_type == out.data_type
        and src.number_format == out.number_format
    )


def verify(old_path, sheet_name, header_row, mappings, new_path, rows_written):
    """Re-read both files and check every migrated cell. Returns the number of cells checked."""
    src_wb = open_workbook(old_path, data_only=True)
    out_wb = open_workbook(new_path)
    try:
        src_ws = pick_sheet(src_wb, old_path, sheet_name)
        out_ws = out_wb.worksheets[0]
        out_ws.reset_dimensions()
        out_rows = out_ws.iter_rows()
        if [c.value for c in next(out_rows, ())] != [m.new_name for m in mappings]:
            raise MigrationError("verification failed: the header row of the new file is not as expected.")

        mismatches = []
        src_rows = src_ws.iter_rows(min_row=header_row + 1)
        for offset, (srow, orow) in enumerate(zip_longest(src_rows, out_rows, fillvalue=())):
            for out_index, m in enumerate(mappings):
                out = at(orow, out_index, EMPTY_CELL)
                src = EMPTY_CELL if m.old_index is None else at(srow, m.old_index, EMPTY_CELL)
                if not same_cell(src, out):
                    old_ref = "(empty column)" if m.old_index is None else cell_ref(m.old_index, header_row + 1 + offset)
                    mismatches.append(f"old {old_ref} -> new {cell_ref(out_index, offset + 2)}")
    finally:
        src_wb.close()
        out_wb.close()

    if mismatches:
        raise MigrationError(f"verification failed, {len(mismatches)} cell(s) differ: {listed(mismatches)}")
    return rows_written * len(mappings)


def partial_path_for(output_path):
    return output_path.with_name(output_path.stem + ".incomplete.xlsx")


def migrate(old_path, helper_path, output_path, sheet_name, header_row, new_sheet):
    mappings = read_mapping(helper_path)

    out_wb = Workbook(write_only=True)
    try:
        out_ws = out_wb.create_sheet(new_sheet)
    except ValueError as exc:
        raise MigrationError(f"--new-sheet '{new_sheet}' is not a valid sheet name ({exc}).") from None
    try:
        rows_written, formula_cells, not_migrated, unnamed = write_new_file(
            old_path, sheet_name, header_row, mappings, out_ws
        )
    except BaseException:
        out_ws.close()  # release openpyxl's temporary row buffer cleanly
        raise

    partial_path = partial_path_for(output_path)
    try:
        out_wb.save(partial_path)
        checked = verify(old_path, sheet_name, header_row, mappings, partial_path, rows_written)
        os.replace(partial_path, output_path)
    except BaseException:
        partial_path.unlink(missing_ok=True)
        raise

    filled = [m for m in mappings if m.old_name]
    empty = [m.new_name for m in mappings if not m.old_name]
    print(f"Old file     : {old_path} (header row {header_row})")
    print(f"Columns      : {len(mappings)} in new file, {len(filled)} filled from old file, {len(empty)} left empty")
    print(f"Rows copied  : {rows_written}")
    print(f"Verified     : {checked} cells match the old file exactly")
    print(f"Saved to     : {output_path}")
    if empty:
        print(f"\nNew columns left empty (no Old Column in the helper): {', '.join(empty)}")
    if not_migrated:
        print(f"\nOld columns NOT migrated (not in the helper): {', '.join(not_migrated)}")
    if unnamed:
        print(f"\nWARNING: old file has data in columns without a header, NOT migrated: {', '.join(unnamed)}")
    if formula_cells:
        print(f"\nNote: {formula_cells} copied cells were formulas in the old file; their results were copied as values.")


def ask(question, default=None, shown_default=None):
    """Ask in the console; Enter alone gives the default.

    Quotes around the answer are removed, so a file dragged into the window works.
    """
    shown = shown_default or default
    answer = input(f"{question}{f' [{shown}]' if shown else ''}: ").strip().strip("\"'").strip()
    return answer or default


def ask_yes_no(question, default):
    while True:
        answer = input(f"{question} ({'Y/n' if default else 'y/N'}): ").strip().lower()
        if not answer:
            return default
        if answer in ("y", "yes", "n", "no"):
            return answer.startswith("y")
        print("   Please type y or n.")


def ask_xlsx_file(question, default=None, shown_default=None):
    while True:
        answer = ask(question, default, shown_default)
        path = Path(answer) if answer else None
        if path is None:
            print("   Please give a file.")
        elif path.suffix.lower() != ".xlsx":
            print("   This is not an .xlsx file. An old .xls file must first be opened in Excel and saved as .xlsx.")
        elif not path.is_file():
            print(f"   File not found: {path}")
        else:
            return path


def ask_for_arguments(default_helper):
    """Build the command-line options by asking step by step. Returns None if the user cancels."""
    print("=" * 70)
    print(" ERP DATA MIGRATION")
    print(" Copies data from an old Excel file into a new Excel file, using the")
    print(" column list in the mapping file (data_migration_helper.xlsx).")
    print()
    print(" Tip: to give a file, drag it from File Explorer into this window and")
    print("      press Enter. Close the Excel files before you start.")
    print("=" * 70)
    print()

    while True:
        old = ask_xlsx_file("1. Old Excel file (the file that has the data)")
        try:
            wb = open_workbook(old)
            sheets = wb.sheetnames
            wb.close()
            break
        except MigrationError as exc:
            print(f"   {exc}")
    argv = [f"--old={old}"]

    sheet = None
    if len(sheets) > 1:
        print("   The old file has these sheets:")
        for number, name in enumerate(sheets, start=1):
            print(f"     {number}. {name}")
        while True:
            answer = ask("   Type the number of the sheet that has the data")
            if answer and answer.isdigit() and 1 <= int(answer) <= len(sheets):
                break
            print(f"   Please type a number from 1 to {len(sheets)}.")
        sheet = sheets[int(answer) - 1]
        argv.append(f"--sheet={sheet}")

    while True:
        header_row = ask("2. Row number that has the column names in the old file", "1")
        if header_row.isdigit() and int(header_row) >= 1:
            break
        print("   Please type a number, for example 1.")
    argv.append(f"--header-row={header_row}")

    helper = ask_xlsx_file(
        "3. Mapping file (press Enter to use the one in this folder)",
        default=str(default_helper),
        shown_default=default_helper.name,
    )
    argv.append(f"--helper={helper}")

    while True:
        answer = ask("4. Name of the new file to create", str(old.with_name(old.stem + "_new.xlsx")))
        output = Path(answer)
        if output.suffix.lower() != ".xlsx":
            output = output.with_name(output.name + ".xlsx")
        if not output.is_absolute() and output.parent == Path("."):
            output = old.parent / output  # a bare name goes next to the old file
        if not output.exists():
            break
        if ask_yes_no(f"   {output} already exists. Replace it?", default=False):
            argv.append("--overwrite")
            break
    argv.append(f"--output={output}")

    print()
    print("Ready to start:")
    print(f"   Old file     : {old}" + (f"  (sheet: {sheet})" if sheet else ""))
    print(f"   Column names : row {header_row}")
    print(f"   Mapping file : {helper}")
    print(f"   New file     : {output}")
    if not ask_yes_no("Start the migration?", default=True):
        return None
    print("\nWorking, please wait...\n")
    return argv


class Parser(argparse.ArgumentParser):
    def error(self, message):
        self.print_usage(sys.stderr)
        self.exit(2, f"ERROR: {message}\nRun '{self.prog} --help' to see what each option means.\n")


HELP_DESCRIPTION = """\
Copies data from an old Excel file into a new Excel file that has different
column names. Which old column goes into which new column is written in the
"Mapping" sheet of data_migration_helper.xlsx.

Easiest way: type only "python main.py", without any options, and it
asks you for everything step by step."""

HELP_EXAMPLES = r"""examples:
  Copy the data from old_data.xlsx into a new file called new_data.xlsx:
    python main.py --old "C:\Data\old_data.xlsx" --output "C:\Data\new_data.xlsx"

  The old file has several sheets; take the data from the sheet "Customers":
    python main.py --old "old_data.xlsx" --sheet "Customers" --output "new_data.xlsx"

  The column names in the old file are in row 3, not in row 1:
    python main.py --old "old_data.xlsx" --header-row 3 --output "new_data.xlsx"

good to know:
  - Give the full path of a file (C:\...) when it is not in the current folder.
  - Put "quotes" around file names and sheet names that contain spaces.
  - Do not create the new file yourself. The script creates it.
  - The old file and the mapping file are only read. They are never changed.
  - Close the Excel files before you run the script."""


def build_parser(default_helper):
    parser = Parser(
        prog="python main.py",
        description=HELP_DESCRIPTION,
        epilog=HELP_EXAMPLES,
        formatter_class=argparse.RawDescriptionHelpFormatter,
        add_help=False,
    )
    must = parser.add_argument_group("you must give")
    must.add_argument("--old", required=True, type=Path, metavar="FILE",
                      help="The old Excel file (.xlsx) that has the data.")
    must.add_argument("--output", required=True, type=Path, metavar="FILE",
                      help="The new Excel file to create, for example new_data.xlsx. "
                           "It must not exist yet (see --overwrite).")
    extra = parser.add_argument_group("extra options (use them only when needed)")
    extra.add_argument("--sheet", metavar="NAME",
                       help="Name of the sheet (tab at the bottom in Excel) in the old file to take the data from. "
                            "Needed only when the old file has more than one sheet.")
    extra.add_argument("--header-row", type=int, default=1, metavar="N",
                       help="Row number in the old file that has the column names. Default: 1.")
    extra.add_argument("--helper", type=Path, default=default_helper, metavar="FILE",
                       help="The mapping file to use. Default: data_migration_helper.xlsx in this folder.")
    extra.add_argument("--new-sheet", default="Sheet1", metavar="NAME",
                       help="Name of the sheet in the new file. Default: Sheet1.")
    extra.add_argument("--overwrite", action="store_true",
                       help="Replace the new file if it already exists. Without this, the script stops "
                            "instead, so no file is replaced by mistake.")
    extra.add_argument("-h", "--help", action="help", help="Show this help and exit.")
    return parser


def main(argv=None):
    if argv is None:
        argv = sys.argv[1:]
    default_helper = Path(__file__).resolve().with_name("data_migration_helper.xlsx")
    if not argv:
        try:
            argv = ask_for_arguments(default_helper)
        except (KeyboardInterrupt, EOFError):
            argv = None
        if argv is None:
            print("\nCancelled. Nothing was created.")
            return 1
    args = build_parser(default_helper).parse_args(argv)

    try:
        for label, path in (("old file", args.old), ("helper file", args.helper)):
            if not path.is_file():
                raise MigrationError(f"{label} not found: {path}")
        if args.output.suffix.lower() != ".xlsx":
            raise MigrationError("--output must end with .xlsx")
        inputs = {args.old.resolve(), args.helper.resolve()}
        if {args.output.resolve(), partial_path_for(args.output).resolve()} & inputs:
            raise MigrationError("--output must be a different file from the old file and the helper.")
        if args.output.exists() and not args.overwrite:
            raise MigrationError(f"{args.output} already exists; use --overwrite to replace it.")
        if not args.output.parent.is_dir():
            raise MigrationError(f"output folder does not exist: {args.output.parent}")
        if args.header_row < 1:
            raise MigrationError("--header-row must be 1 or more.")
        migrate(args.old, args.helper, args.output, args.sheet, args.header_row, args.new_sheet)
    except MigrationError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    except PermissionError as exc:
        print(f"ERROR: {exc}. Close the file if it is open in Excel and try again.", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
